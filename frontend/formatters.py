"""Text clean-up and export helpers: TXT / DOCX / PDF / HTML preview."""
import html
import io
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from config import FOOTER_TEXT, LOGO_PATH  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402
from fpdf import FPDF  # noqa: E402

_REPLACEMENTS = {
    "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
    "\u2013": "-", "\u2014": "-", "\u2022": "-", "\u2026": "...",
    "\u20b9": "Rs. ", "\u00a0": " ",
}


def sanitize_text(text: str) -> str:
    """Remove markdown symbols, typographic quotes and characters PDF fonts can't print."""
    for old, new in _REPLACEMENTS.items():
        text = text.replace(old, new)
    text = re.sub(r"\*\*|__", "", text)                       # bold markers
    text = re.sub(r"^\s{0,3}#{1,6}\s*", "", text, flags=re.M)  # markdown headings
    text = re.sub(r"^\s*\*\s+", "- ", text, flags=re.M)        # '* item' -> '- item'
    text = text.encode("latin-1", "ignore").decode("latin-1")
    return text.strip()


def is_heading(line: str) -> bool:
    s = line.strip()
    if not s or len(s) > 90:
        return False
    if re.match(r"^\d+\.\s+\S", s) and (s.endswith(":") or len(s) < 60):
        return True
    return s.isupper() and len(s) > 3


def is_bullet(line: str) -> bool:
    return bool(re.match(r"^\s*[-*]\s+\S", line))


def strip_bullet(line: str) -> str:
    return re.sub(r"^\s*[-*]\s+", "", line).strip()


def split_terms(terms: str):
    return [t.strip() for t in (terms or "").split(";") if t.strip()]


# ---------------------------------------------------------------- HTML preview
def format_html_preview(text: str) -> str:
    out = []
    for line in text.splitlines():
        if not line.strip():
            out.append("<br>")
        elif is_heading(line):
            out.append(f"<h4 style='margin:14px 0 4px;color:#172b4d'>{html.escape(line.strip())}</h4>")
        elif is_bullet(line):
            out.append(f"<div style='margin-left:18px'>&bull; {html.escape(strip_bullet(line))}</div>")
        else:
            out.append(f"<p style='margin:4px 0'>{html.escape(line.strip())}</p>")
    return "".join(out)


# ------------------------------------------------------------------------ DOCX
def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    if os.path.exists(LOGO_PATH):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(LOGO_PATH, width=Inches(2.2))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(doc_type.upper())
    r.bold = True
    r.font.size = Pt(16)

    term_list = split_terms(terms)
    if term_list:
        doc.add_paragraph().add_run("Summary of Key Terms").bold = True
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term / Condition"
        for c in table.rows[0].cells:
            for run in c.paragraphs[0].runs:
                run.bold = True
        for i, t in enumerate(term_list, 1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = t
        doc.add_paragraph()

    for line in text.splitlines():
        if not line.strip():
            continue
        if is_heading(line):
            doc.add_paragraph().add_run(line.strip()).bold = True
        elif is_bullet(line):
            doc.add_paragraph(strip_bullet(line), style="List Bullet")
        else:
            doc.add_paragraph(line.strip())

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = FOOTER_TEXT
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------------- PDF
class _LegalPDF(FPDF):
    def __init__(self, doc_type):
        super().__init__()
        self.doc_type = doc_type
        self.set_margins(20, 38, 20)
        self.set_auto_page_break(auto=True, margin=22)
        self.alias_nb_pages()

    def header(self):
        if os.path.exists(LOGO_PATH):
            self.image(LOGO_PATH, x=(self.w - 38) / 2, y=8, w=38)
        self.set_y(26)
        self.set_font("Helvetica", "B", 12)
        self.cell(0, 8, self.doc_type, align="C", new_x="LMARGIN", new_y="NEXT")
        self.set_y(38)

    def footer(self):
        self.set_y(-16)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 6, FOOTER_TEXT, align="C", new_x="LMARGIN", new_y="NEXT")
        self.cell(0, 6, f"Page {self.page_no()} / {{nb}}", align="C")


def format_pdf(text: str, doc_type: str, terms: str = "") -> bytes:
    doc_type = sanitize_text(doc_type)
    pdf = _LegalPDF(doc_type)
    pdf.add_page()

    def write(line, style="", size=11, indent=0):
        pdf.set_font("Helvetica", style, size)
        pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(0, 6, line, new_x="LMARGIN", new_y="NEXT")

    term_list = split_terms(terms)
    if term_list:
        write("Summary of Key Terms", "B", 11)
        for t in term_list:
            write("-  " + sanitize_text(t), indent=4)
        pdf.ln(3)

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            pdf.ln(3)
        elif is_heading(line):
            pdf.ln(2)
            write(line, "B", 11)
        elif is_bullet(raw):
            write("-  " + strip_bullet(raw), indent=4)
        else:
            write(line)
    return bytes(pdf.output())
