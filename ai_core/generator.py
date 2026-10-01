"""Formatting helpers: sanitize, HTML preview, DOCX and PDF export."""
import html
import io
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF

from config import FOOTER_TEXT, LOGO_PATH

_REPLACEMENTS = {
    "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
    "\u2013": "-", "\u2014": "-", "\u2026": "...", "\u00a0": " ", "\u2022": "-",
}


def sanitize_text(text: str) -> str:
    """Replace typographic characters and strip control characters."""
    for src, dst in _REPLACEMENTS.items():
        text = text.replace(src, dst)
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)


def _clean_line(line: str) -> str:
    line = re.sub(r"^\s*#+\s*", "", line)
    return line.replace("**", "").replace("__", "").rstrip()


def _is_heading(line: str) -> bool:
    s = line.strip()
    if not s:
        return False
    return bool(
        re.match(r"^\d+\.\s+\S", s) and len(s) < 70
        or (s.endswith(":") and len(s) < 60)
        or (s.isupper() and len(s) < 60)
    )


def split_terms(terms: str):
    return [t.strip() for t in (terms or "").split(";") if t.strip()]


# ---------------------------------------------------------------- HTML preview
def format_html_preview(text: str) -> str:
    blocks = []
    for raw in text.splitlines():
        line = _clean_line(raw)
        if not line.strip():
            continue
        safe = html.escape(line)
        if raw.lstrip().startswith("#"):
            blocks.append(f"<h3 style='margin:14px 0 6px'>{safe}</h3>")
        elif _is_heading(line):
            blocks.append(f"<p style='margin:14px 0 4px'><b>{safe}</b></p>")
        else:
            blocks.append(f"<p style='margin:4px 0;line-height:1.55'>{safe}</p>")
    return "".join(blocks)


# ------------------------------------------------------------------------ DOCX
def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    if os.path.exists(LOGO_PATH):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(LOGO_PATH, width=Inches(2.2))

    title = doc.add_heading(doc_type or "Legal Document", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for raw in text.splitlines():
        line = _clean_line(raw)
        if not line.strip():
            continue
        p = doc.add_paragraph()
        run = p.add_run(line)
        if _is_heading(line):
            run.bold = True
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    items = split_terms(terms)
    if items:
        doc.add_paragraph().add_run("Summary of Key Terms").bold = True
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term"
        for i, item in enumerate(items, 1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = item

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = FOOTER_TEXT
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------------- PDF
class _LegalPDF(FPDF):
    def __init__(self, doc_type: str):
        super().__init__()
        self.doc_type = doc_type

    def header(self):
        if os.path.exists(LOGO_PATH):
            self.image(LOGO_PATH, x=(self.w - 45) / 2, y=8, w=45)
        self.set_y(28)
        self.set_font("Helvetica", "B", 12)
        self.cell(0, 8, _latin(self.doc_type), align="C")
        self.ln(12)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 10, _latin(FOOTER_TEXT), align="C")


def _latin(s: str) -> str:
    return sanitize_text(s).encode("latin-1", "replace").decode("latin-1")


def format_pdf(text: str, doc_type: str, terms: str = "") -> bytes:
    pdf = _LegalPDF(doc_type or "Legal Document")
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    for raw in text.splitlines():
        line = _latin(_clean_line(raw))
        if not line.strip():
            continue
        pdf.set_x(pdf.l_margin)
        if _is_heading(line):
            pdf.set_font("Helvetica", "B", 11)
            pdf.multi_cell(0, 7, line)
        else:
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(0, 6, line)

    items = split_terms(terms)
    if items:
        pdf.ln(4)
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "B", 11)
        pdf.multi_cell(0, 7, "Summary of Key Terms")
        pdf.set_font("Helvetica", "", 10)
        for item in items:
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(0, 6, "- " + _latin(item))

    out = pdf.output(dest="S")
    return out.encode("latin-1") if isinstance(out, str) else bytes(out)
